#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Minimal Markdown -> DOCX converter tuned for the PocketLedger plan document.

Supports: h1-h4, tables, bullet lists, numbered lists, task checkboxes,
fenced code blocks, blockquotes, **bold** and `code` inline spans.

Usage: python md2docx.py <input.md> <output.docx> [--title "文档标题"]
"""
from __future__ import annotations

import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

CJK = "微软雅黑"
BODY_FONT = "Segoe UI"
MONO_FONT = "Consolas"
ACCENT = RGBColor(0x1F, 0x3A, 0x5F)

INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")
MD_LINK = re.compile(r"^\[([^\]]+)\]\(([^)]+)\)$")
SEP_CELL = re.compile(r"^:?-{2,}:?$")


def set_run_fonts(run, size=None, bold=None, italic=None, mono=False, color=None):
    name = MONO_FONT if mono else BODY_FONT
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:eastAsia"), CJK)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_inline(par, text, size=None, mono=False, color=None):
    for part in INLINE.split(text):
        if not part:
            continue
        link = MD_LINK.match(part)
        if link:
            label, url = link.group(1), link.group(2)
            set_run_fonts(par.add_run(label), size=size, mono=mono, color=color)
            small = (size - 1.5) if size else None
            set_run_fonts(
                par.add_run(f"（{url}）"), size=small, color=RGBColor(0x55, 0x55, 0x55)
            )
        elif part.startswith("**") and part.endswith("**") and len(part) > 4:
            set_run_fonts(par.add_run(part[2:-2]), size=size, bold=True, mono=mono, color=color)
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            set_run_fonts(par.add_run(part[1:-1]), size=size, mono=True, color=color)
        else:
            set_run_fonts(par.add_run(part), size=size, mono=mono, color=color)


def style_heading(styles, name, size):
    st = styles[name]
    st.font.name = BODY_FONT
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = ACCENT
    rPr = st.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), BODY_FONT)
    rFonts.set(qn("w:hAnsi"), BODY_FONT)
    rFonts.set(qn("w:eastAsia"), CJK)
    st.paragraph_format.space_before = Pt(12 if size >= 14 else 8)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.keep_with_next = True


def setup(doc, title):
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(1.8)
    sec.right_margin = Cm(1.8)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(1.8)

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), BODY_FONT)
    rFonts.set(qn("w:hAnsi"), BODY_FONT)
    rFonts.set(qn("w:eastAsia"), CJK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.28

    for name, size in (
        ("Title", 22),
        ("Heading 1", 16),
        ("Heading 2", 13.5),
        ("Heading 3", 12),
        ("Heading 4", 11),
    ):
        style_heading(doc.styles, name, size)

    footer = sec.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_fonts(p.add_run(f"{title}   ·   第 "), size=8, color=RGBColor(0x80, 0x80, 0x80))
    run = p.add_run()
    for el, attrs, text in (
        ("w:fldChar", {"w:fldCharType": "begin"}, None),
        ("w:instrText", {"xml:space": "preserve"}, " PAGE "),
        ("w:fldChar", {"w:fldCharType": "end"}, None),
    ):
        e = OxmlElement(el)
        for k, v in attrs.items():
            e.set(qn(k) if k != "xml:space" else "{http://www.w3.org/XML/1998/namespace}space", v)
        if text:
            e.text = text
        run._element.append(e)
    set_run_fonts(run, size=8, color=RGBColor(0x80, 0x80, 0x80))
    set_run_fonts(p.add_run(" 页"), size=8, color=RGBColor(0x80, 0x80, 0x80))


def add_table(doc, rows):
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    table = doc.add_table(rows=1, cols=width)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    for i, cell_text in enumerate(rows[0]):
        cell = table.rows[0].cells[i]
        par = cell.paragraphs[0]
        par.paragraph_format.space_after = Pt(2)
        par.paragraph_format.line_spacing = 1.1
        add_inline(par, cell_text, size=9)
        for run in par.runs:
            run.bold = True
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), "EAF0F7")
        cell._tc.get_or_add_tcPr().append(shd)

    for row in rows[1:]:
        cells = table.add_row().cells
        for i, cell_text in enumerate(row):
            par = cells[i].paragraphs[0]
            par.paragraph_format.space_after = Pt(2)
            par.paragraph_format.line_spacing = 1.1
            add_inline(par, cell_text, size=9)

    trPr = table.rows[0]._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    trPr.append(header)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def parse_table_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def convert(md_text, out_path, title):
    doc = Document()
    setup(doc, title)

    lines = md_text.splitlines()
    i = 0
    in_code = False
    code_buf: list[str] = []
    first_heading_done = False

    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()

        if line.strip().startswith("```"):
            if in_code:
                for cl in code_buf:
                    par = doc.add_paragraph()
                    pf = par.paragraph_format
                    pf.space_after = Pt(0)
                    pf.space_before = Pt(0)
                    pf.line_spacing = 1.0
                    pf.left_indent = Cm(0.3)
                    set_run_fonts(par.add_run(cl if cl.strip() else " "), size=8.5, mono=True)
                code_buf = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue

        if in_code:
            code_buf.append(raw)
            i += 1
            continue

        stripped = line.strip()

        if not stripped or stripped.startswith("<!--"):
            i += 1
            continue

        if set(stripped) <= {"-", " "} and len(stripped) >= 3:
            i += 1
            continue

        if stripped.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            rows = []
            for k, bl in enumerate(block):
                cells = parse_table_row(bl)
                if k == 1 and all(SEP_CELL.match(c) for c in cells if c):
                    continue
                rows.append(cells)
            if rows:
                add_table(doc, rows)
            continue

        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            if level == 1 and not first_heading_done:
                par = doc.add_paragraph(style="Title")
                par.paragraph_format.space_after = Pt(10)
                add_inline(par, text, size=22, color=ACCENT)
                for run in par.runs:
                    run.bold = True
                first_heading_done = True
            else:
                par = doc.add_paragraph(style=f"Heading {min(level - 1, 4)}")
                add_inline(par, text)
            i += 1
            continue

        if stripped.startswith(">"):
            text = stripped.lstrip(">").strip()
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.6)
            par.paragraph_format.space_before = Pt(4)
            add_inline(par, text, size=10, color=RGBColor(0x44, 0x44, 0x44))
            for run in par.runs:
                run.italic = True
            i += 1
            continue

        m = re.match(r"^(\s*)[-*]\s+(.*)$", line)
        if m:
            text = m.group(2).strip()
            if text.startswith("[ ]"):
                text = "☐ " + text[3:].strip()
            elif text.lower().startswith("[x]"):
                text = "☑ " + text[3:].strip()
            par = doc.add_paragraph(style="List Bullet")
            par.paragraph_format.space_after = Pt(3)
            par.paragraph_format.line_spacing = 1.2
            add_inline(par, text)
            i += 1
            continue

        m = re.match(r"^\s*(\d+)\.\s+(.*)$", line)
        if m:
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.6)
            par.paragraph_format.space_after = Pt(3)
            add_inline(par, f"{m.group(1)}. {m.group(2).strip()}")
            i += 1
            continue

        par = doc.add_paragraph()
        add_inline(par, stripped)
        i += 1

    doc.save(out_path)
    print(f"OK -> {out_path}")


def main():
    args = [a for a in sys.argv[1:]]
    title = "PocketLedger 技术方案"
    if "--title" in args:
        k = args.index("--title")
        title = args[k + 1]
        del args[k : k + 2]
    src, dst = args[0], args[1]
    with open(src, encoding="utf-8") as fh:
        convert(fh.read(), dst, title)


if __name__ == "__main__":
    main()
